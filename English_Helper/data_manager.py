"""
英文單字訓練系統 - 數據管理層
處理詞庫和試題數據的讀寫
"""

import json
from pathlib import Path
from typing import List, Dict, Optional
import random
from datetime import datetime
from utils.app_logger import get_app_logger


LOGGER = get_app_logger("english_helper.data")


class DataManager:
    """數據管理層 - 讀寫詞庫和試題"""
    
    def __init__(self):
        self.vocabulary_file = Path("vocabulary_core.json")
        self.custom_vocab_file = Path("vocabulary_custom.json")
        self.exam_file = Path("exam_questions.xlsx")
        self.records_file = Path("training_records.json")
        self.daily_vocab_stats_file = Path("daily_vocab_stats.json")
        
        # 緩存
        self._vocabulary = None
        self._custom_vocabulary = None
        self._exam_questions = None
        self._daily_vocab_stats = None
        
        self._load_or_create_custom_vocab()
        self._load_or_create_records()
        self._load_or_create_daily_vocab_stats()
        LOGGER.info("DataManager initialized")
    
    def _load_or_create_custom_vocab(self):
        """加載或創建自訂詞庫"""
        if self.custom_vocab_file.exists():
            with open(self.custom_vocab_file, 'r', encoding='utf-8') as f:
                self._custom_vocabulary = json.load(f).get('custom_words', [])
            LOGGER.info("Loaded custom vocabulary count=%s", len(self._custom_vocabulary))
        else:
            self._custom_vocabulary = []
            self._save_custom_vocab()
            LOGGER.info("Created new custom vocabulary file")
    
    def _load_or_create_records(self):
        """加載或創建訓練紀錄"""
        if self.records_file.exists():
            with open(self.records_file, 'r', encoding='utf-8') as f:
                self.records = json.load(f).get('records', [])
            LOGGER.info("Loaded training records count=%s", len(self.records))
        else:
            self.records = []
            self._save_records()
            LOGGER.info("Created new training records file")

    def _load_or_create_daily_vocab_stats(self):
        """加載或創建今日單字訓練統計"""
        today = datetime.now().date().isoformat()

        if self.daily_vocab_stats_file.exists():
            with open(self.daily_vocab_stats_file, 'r', encoding='utf-8') as f:
                data = json.load(f)

            if data.get('date') == today:
                self._daily_vocab_stats = {
                    'date': today,
                    'answered': int(data.get('answered', 0)),
                    'correct': int(data.get('correct', 0)),
                    'wrong': int(data.get('wrong', 0)),
                }
                LOGGER.info(
                    "Loaded daily vocab stats date=%s answered=%s correct=%s wrong=%s",
                    today,
                    self._daily_vocab_stats['answered'],
                    self._daily_vocab_stats['correct'],
                    self._daily_vocab_stats['wrong'],
                )
                return

        self._daily_vocab_stats = {
            'date': today,
            'answered': 0,
            'correct': 0,
            'wrong': 0,
        }
        self._save_daily_vocab_stats()
        LOGGER.info("Created new daily vocab stats file date=%s", today)
    
    def get_vocabulary(self) -> List[Dict]:
        """取得核心詞庫"""
        if self._vocabulary is None:
            with open(self.vocabulary_file, 'r', encoding='utf-8') as f:
                data = json.load(f)
                self._vocabulary = data.get('words', [])
        return self._vocabulary
    
    def get_custom_vocabulary(self) -> List[Dict]:
        """取得自訂詞庫"""
        return self._custom_vocabulary
    
    def get_all_vocabulary(self) -> List[Dict]:
        """取得所有詞庫（核心+自訂）"""
        return self.get_vocabulary() + self.get_custom_vocabulary()
    
    def add_custom_word(self, english: str, chinese: str, pos: str = "", phonetic: str = "", example: str = ""):
        """添加自訂單字"""
        new_word = {
            'english': english.lower(),
            'chinese': chinese,
            'part_of_speech': pos,
        }
        
        self._custom_vocabulary.append(new_word)
        self._save_custom_vocab()
        LOGGER.info("Added custom word english=%s", new_word['english'])
        return new_word

    def update_custom_word(self, word_index: int, english: str, chinese: str, pos: str = ""):
        """更新自訂單字"""
        if 0 <= word_index < len(self._custom_vocabulary):
            self._custom_vocabulary[word_index] = {
                'english': english.lower(),
                'chinese': chinese,
                'part_of_speech': pos,
            }
            self._save_custom_vocab()
            LOGGER.info("Updated custom word english=%s index=%s", english.lower(), word_index)
            return self._custom_vocabulary[word_index]

        LOGGER.warning("Custom word not found for update index=%s", word_index)
        return None

    def replace_custom_vocabulary(self, words: List[Dict]):
        """整批覆寫自訂詞庫"""
        cleaned_words = []
        for word in words:
            cleaned_words.append({
                'english': word.get('english', '').strip().lower(),
                'chinese': word.get('chinese', '').strip(),
                'part_of_speech': word.get('part_of_speech', '').strip(),
            })

        self._custom_vocabulary = cleaned_words
        self._save_custom_vocab()
        LOGGER.info("Replaced custom vocabulary count=%s", len(cleaned_words))
    
    def delete_custom_word(self, word_index: int):
        """刪除自訂單字"""
        if 0 <= word_index < len(self._custom_vocabulary):
            del self._custom_vocabulary[word_index]
        else:
            LOGGER.warning("Custom word not found for delete index=%s", word_index)
            return
        self._save_custom_vocab()
        LOGGER.info("Deleted custom word index=%s", word_index)
    
    def _save_custom_vocab(self):
        """保存自訂詞庫"""
        data = {'custom_words': self._custom_vocabulary}
        with open(self.custom_vocab_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        LOGGER.info("Saved custom vocabulary count=%s", len(self._custom_vocabulary))
    
    def _save_records(self):
        """保存訓練紀錄"""
        data = {'records': self.records}
        with open(self.records_file, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        LOGGER.info("Saved training records count=%s", len(self.records))

    def _save_daily_vocab_stats(self):
        """保存今日單字訓練統計"""
        if self._daily_vocab_stats is None:
            return

        with open(self.daily_vocab_stats_file, 'w', encoding='utf-8') as f:
            json.dump(self._daily_vocab_stats, f, ensure_ascii=False, indent=2)
        LOGGER.info(
            "Saved daily vocab stats date=%s answered=%s correct=%s wrong=%s",
            self._daily_vocab_stats['date'],
            self._daily_vocab_stats['answered'],
            self._daily_vocab_stats['correct'],
            self._daily_vocab_stats['wrong'],
        )
    
    def add_training_record(self, session_data: Dict):
        """添加訓練紀錄"""
        self.records.append(session_data)
        self._save_records()
        LOGGER.info("Appended training record total=%s", len(self.records))

    def get_training_records(self) -> List[Dict]:
        """取得所有測驗紀錄（新到舊）"""
        return list(reversed(self.records))

    def normalize_training_record(self, record: Dict) -> Dict:
        """將舊版與新版紀錄統一成 UI 可讀格式"""
        total_questions = int(record.get('total_questions', 0) or 0)
        correct_answers = int(record.get('correct_answers', 0) or 0)
        accuracy = float(record.get('accuracy', 0) or 0)
        spent_time = float(record.get('spent_time', 0) or 0)

        completed_at = (
            record.get('completed_at')
            or record.get('timestamp')
            or ''
        )
        record_type = record.get('record_type') or '測驗'
        if record_type == '測驗':
            record_type = '隨機測驗'

        item_name = record.get('item_name') or record.get('mode_name')
        if not item_name:
            item_name = '隨機測驗' if record_type == '隨機測驗' else '未知測驗'

        return {
            'completed_at': completed_at,
            'record_type': record_type,
            'item_name': item_name,
            'total_questions': total_questions,
            'correct_answers': correct_answers,
            'wrong_answers': max(total_questions - correct_answers, 0),
            'accuracy': accuracy,
            'spent_time': spent_time,
        }

    def record_daily_vocab_answer(self, is_correct: bool):
        """記錄今日單字作答統計"""
        today = datetime.now().date().isoformat()

        if self._daily_vocab_stats is None or self._daily_vocab_stats.get('date') != today:
            self._daily_vocab_stats = {
                'date': today,
                'answered': 0,
                'correct': 0,
                'wrong': 0,
            }

        self._daily_vocab_stats['answered'] += 1
        if is_correct:
            self._daily_vocab_stats['correct'] += 1
        else:
            self._daily_vocab_stats['wrong'] += 1

        self._save_daily_vocab_stats()
        return self._daily_vocab_stats

    def get_daily_vocab_stats(self) -> Dict:
        """取得今日單字訓練統計"""
        if self._daily_vocab_stats is None:
            self._load_or_create_daily_vocab_stats()

        return self._daily_vocab_stats or {
            'date': datetime.now().date().isoformat(),
            'answered': 0,
            'correct': 0,
            'wrong': 0,
        }
    
    def get_recent_records(self, count: int = 10) -> List[Dict]:
        """獲取最近的訓練紀錄"""
        return self.records[-count:]
    
    def load_exam_questions(self):
        """從Excel加載歷屆試題"""
        try:
            from openpyxl import load_workbook
            
            wb = load_workbook(self.exam_file)
            ws = wb.active
            
            questions = []
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
                if not row[0]:
                    continue
                
                year, question_type, number, question_text = row[0], row[1], row[2], row[3]
                options = list(row[4:8])  # A, B, C, D
                answer = row[8] if len(row) > 8 else None
                if answer:
                    answer = str(answer).strip().upper()
                    if answer not in {'A', 'B', 'C', 'D'}:
                        answer = None
                
                questions.append({
                    'id': len(questions) + 1,
                    'year': year,
                    'type': question_type,
                    'number': number,
                    'question': question_text,
                    'options': {
                        'A': options[0] if len(options) > 0 else '',
                        'B': options[1] if len(options) > 1 else '',
                        'C': options[2] if len(options) > 2 else '',
                        'D': options[3] if len(options) > 3 else '',
                    },
                    'correct_option': answer,
                })
            
            self._exam_questions = questions
            LOGGER.info("Loaded exam questions from excel count=%s", len(questions))
            return questions
            
        except Exception as e:
            LOGGER.exception("Failed loading exam questions: %s", e)
            print(f"❌ 加載試題失敗: {e}")
            return []
    
    def get_exam_questions(self) -> List[Dict]:
        """取得歷屆試題"""
        if self._exam_questions is None:
            self.load_exam_questions()
        return self._exam_questions or []
    
    def filter_exam_questions(self, year: Optional[int] = None, question_type: Optional[str] = None) -> List[Dict]:
        """篩選試題"""
        questions = self.get_exam_questions()
        
        if year:
            questions = [q for q in questions if q['year'] == year]
        
        if question_type:
            questions = [q for q in questions if q['type'] == question_type]
        
        return questions


# 全局實例
_data_manager = None


def get_data_manager() -> DataManager:
    """獲取全局數據管理器實例"""
    global _data_manager
    if _data_manager is None:
        _data_manager = DataManager()
    return _data_manager
